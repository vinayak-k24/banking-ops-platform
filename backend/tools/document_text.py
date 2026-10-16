"""Text extraction for uploaded migration documents. Returns one string per (pseudo) page."""

import io

PSEUDO_PAGE_CHARS = 3000


class DocumentTooLarge(Exception):
    pass


def _paginate(text: str) -> list[str]:
    """Split non-paged formats into ~3000-character pages on paragraph boundaries."""
    pages, current = [], ""
    for para in text.split("\n\n"):
        if current and len(current) + len(para) > PSEUDO_PAGE_CHARS:
            pages.append(current)
            current = ""
        current += para + "\n\n"
    if current.strip():
        pages.append(current)
    return pages


def extract_pages(data: bytes, ext: str, max_pages: int) -> list[str]:
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        if len(reader.pages) > max_pages:
            raise DocumentTooLarge(f"Document has {len(reader.pages)} pages (limit {max_pages})")
        pages = [page.extract_text() or "" for page in reader.pages]
    elif ext == ".docx":
        from docx import Document

        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            parts += [" | ".join(cell.text for cell in row.cells) for row in table.rows]
        pages = _paginate("\n\n".join(parts))
    else:
        pages = _paginate(data.decode("utf-8", errors="replace"))

    if len(pages) > max_pages:
        raise DocumentTooLarge(f"Document has {len(pages)} pages (limit {max_pages})")
    return pages
